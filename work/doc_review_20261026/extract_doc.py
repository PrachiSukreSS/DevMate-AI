from docx import Document
from docx.oxml.ns import qn
from zipfile import ZipFile
from pathlib import Path

source = Path(r"C:\Users\prach\Downloads\GENAI PROJECT22.docx")
out = Path(r"C:\Users\prach\Documents\Codex\2026-10-22\x20-a\devmate-ai\work\doc_review_20261026\document_text.txt")
doc = Document(source)
lines = []
for child in doc.element.body.iterchildren():
    if child.tag == qn("w:p"):
        text = "".join(child.itertext())
        text = " ".join(text.split())
        if text:
            lines.append(text)
    elif child.tag == qn("w:tbl"):
        lines.append("[TABLE]")
        for row in child.findall(".//" + qn("w:tr")):
            cells=[]
            for cell in row.findall("./" + qn("w:tc")):
                cells.append(" ".join("".join(cell.itertext()).split()))
            lines.append(" | ".join(cells))
        lines.append("[/TABLE]")
with out.open("w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print(f"Extracted {len(lines)} blocks to {out}")
print(f"Paragraphs={len(doc.paragraphs)} Tables={len(doc.tables)} InlineImages={len(doc.inline_shapes)}")
with ZipFile(source) as z:
    media=[n for n in z.namelist() if n.startswith("word/media/")]
    print("Media files:", ", ".join(media))
